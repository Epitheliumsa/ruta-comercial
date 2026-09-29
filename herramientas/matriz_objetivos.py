#!/usr/bin/env python3
"""Convierte la matriz de objetivos y subcategorías (Excel) en objetivos.js, que es lo que lee la app.
Hoja "Tipo de visita": en qué tipos de visita sale cada clasificación de cliente.

Uso:  python3 herramientas/matriz_objetivos.py [ruta del Excel]
      (por defecto: datos/Matriz_App.xlsx, la matriz completa de la app)

Reglas del archivo:
- Hojas "Visitas" y "Trabajo interno": fila con Objetivo (columna A) = objetivo; ✓ o X en los tipos donde sale.
  Filas de abajo con Subcategoría (columna B) = sus subcategorías; X en los tipos donde salen.
  Subcategoría "Variable" = se carga cada mes (hoja "Mensual" o pantalla de jefes en la app).
- Hoja "Mensual": Mes (AAAA-MM) | Objetivo | Subcategoría, para los objetivos variables.
- Orden: objetivos y subcategorías en orden alfabético, salvo Planeación Mes (PRIMEROS) y las subcategorías de ORDEN_SUBS.
"""
import json, sys, unicodedata
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'datos' / 'Matriz_App.xlsx'
SALIDA = RAIZ / 'objetivos.js'
PRIMEROS = {'Planeación Mes': ['Visiplan', 'Diagnóstico de Zona', 'Plan de Acción', 'Plan de Trabajo Diario']}
# Subcategorías con orden fijo (las demás van en orden alfabético). Las que no estén en la lista van al final.
ORDEN_SUBS = {
    'Colocación': ['Producto terminado', 'Magistral individual', 'Magistral de pedido', 'Producto nuevo'],
    'Desarrollo Productos': ['Fórmula magistral nueva', 'Ajuste de fórmula', 'Muestra de desarrollo'],
    'Devoluciones - PQR': ['Devolución', 'Queja', 'Reclamo', 'Reacondicionamiento', 'Sugerencia'],
    'Mapa del Cliente - Ampliación Portafolio': ['Productos nuevos', 'Productos foco', 'Productos transición', 'Portafolio actual'],
    'Productos Nuevos': ['Presentación del producto', 'Entrega de muestra', 'Material de apoyo', 'Codificación'],
}

# Orden que el usuario numera en la columna A de cada subcategoría (manda sobre ORDEN_SUBS y el alfabético)
ORDEN_NUM = {}
def ordenar_subs(obj, lista):
    fijo = [clave(x) for x in ORDEN_SUBS.get(obj, [])]
    def llave(x):
        k = clave(x)
        return (ORDEN_NUM.get((obj, k), 999), fijo.index(k) if k in fijo else 999, k)
    return sorted(lista, key=llave)
# Nombre de la columna en el Excel -> clave del tipo en la app
TIPOS = {
    'Visita Médica': 'Visita Médica', 'Visita Comercial': 'Visita Comercial', 'Punto de Venta': 'Punto de Venta',
    # Visita Médica a clientes con clasificación 20 y 21 (médico que también compra): tiene sus propios objetivos
    'Visita Médica Comercial (clasificación 20 y 21)': 'medcom:Visita Médica',
    'Contacto nuevo · Visita Médica': 'nuevo:Visita Médica', 'Contacto nuevo · Visita Comercial': 'nuevo:Visita Comercial',
    'Contacto nuevo · Punto de Venta': 'nuevo:Punto de Venta',
    'Trabajo Administrativo Oficina': 'Trabajo Administrativo Oficina',
    'Trabajo Administrativo Fuera de la Oficina': 'Trabajo Administrativo Fuera de la Oficina',
    'Planeación Mes': 'Planeación Mes',
}

def clave(t):
    return unicodedata.normalize('NFD', t.lower()).encode('ascii', 'ignore').decode()

def limpio(v):
    return ' '.join(str(v).split()) if v not in (None, '') else ''

# Objetivos y subcategorías en nombre propio (como NOMPROPIO de Excel), con los conectores en minúscula
# ("Chequeo de Precios") y las siglas como vienen (PQR, SPA)
MENORES = {'de', 'del', 'la', 'las', 'el', 'los', 'y', 'e', 'o', 'u', 'a', 'al', 'en', 'con', 'para', 'por', 'sin', 'su', 'sus', 'un', 'una'}
def nomprop(t):
    palabras = t.split(' ')
    return ' '.join(p.lower() if i and p.lower() in MENORES else p if len(p) > 1 and p.isupper() else p[:1].upper() + p[1:]
                    for i, p in enumerate(palabras))

def ordenar(tipo, lista):
    primeros = [o for o in PRIMEROS.get(tipo, []) if o in lista]
    return primeros + sorted([o for o in lista if o not in primeros], key=clave)

libro = openpyxl.load_workbook(ENTRADA, data_only=True)
objetivos, subcategorias, variables, avisos = {}, {}, [], []
for hoja in ('Visitas', 'Trabajo interno'):
    ws = libro[hoja]
    cab = [limpio(c.value) for c in ws[4]]
    cols = {i: TIPOS[n] for i, n in enumerate(cab) if n in TIPOS}
    for n in cab[2:]:
        if n and n not in TIPOS:
            avisos.append(f'{hoja}: columna "{n}" no es un tipo conocido (se ignora)')
    actual = None
    for fila in ws.iter_rows(min_row=5, values_only=True):
        obj, sub = limpio(fila[0]), limpio(fila[1]) if len(fila) > 1 else ''
        if obj.startswith('✓ en la'):
            break
        # Un número en la columna A de una subcategoría es su orden dentro del objetivo
        if obj.replace('.0', '').isdigit() and sub:
            if actual:
                ORDEN_NUM[(actual, clave(nomprop(sub)))] = int(float(obj))
            obj = ''
        obj = nomprop(obj) if obj else obj
        sub = sub if not sub or sub.lower().startswith('variable') else nomprop(sub)
        if obj:
            actual = obj
            for i, t in cols.items():
                if limpio(fila[i]):
                    objetivos.setdefault(t, []).append(obj)
        elif sub and actual:
            if sub.lower().startswith('variable'):
                if actual not in variables:
                    variables.append(actual)
                continue
            for i, t in cols.items():
                if limpio(fila[i]):
                    if actual not in objetivos.get(t, []):
                        avisos.append(f'{hoja}: "{sub}" tiene X en {t} pero "{actual}" no sale en ese tipo')
                    subcategorias.setdefault(t, {}).setdefault(actual, []).append(sub)

objetivos = {t: ordenar(t, l) for t, l in objetivos.items()}
subcategorias = {t: {o: ordenar_subs(o, set(s)) for o, s in d.items()} for t, d in subcategorias.items()}

mensual = {}
if 'Mensual' in libro.sheetnames:
    for fila in libro['Mensual'].iter_rows(min_row=5, values_only=True):
        mes, obj, sub = (limpio(v) for v in (list(fila) + [None] * 3)[:3])
        if mes and obj and sub:
            mes = mes[:7]
            mensual.setdefault(mes, {}).setdefault(obj, []).append(sub)

# Clasificación del cliente -> tipos de visita donde sale (X en la matriz). Con Visita Médica y Visita Comercial
# a la vez, al programar el vendedor marca una, otra o ambas.
por_clasificacion = {}
if 'Tipo de visita' in libro.sheetnames:
    ws = libro['Tipo de visita']
    cab = [limpio(c.value) for c in ws[4]]
    cols = {i: n for i, n in enumerate(cab) if n in ('Visita Médica', 'Visita Comercial', 'Punto de Venta')}
    for fila in ws.iter_rows(min_row=5, values_only=True):
        cl = limpio(fila[0])
        if not cl.isdigit():
            continue
        tipos = [n for i, n in cols.items() if limpio(fila[i])]
        if tipos:
            por_clasificacion[cl] = tipos
        else:
            avisos.append(f'Clasificación {cl} no tiene X en ningún tipo de visita (sale en todos)')

orden_tipos = [t for t in TIPOS.values() if t in objetivos]
datos = {'objetivos': {t: objetivos[t] for t in orden_tipos}, 'subcategorias': subcategorias,
         'variables': variables, 'mensual': mensual, 'tiposPorClasificacion': por_clasificacion}
SALIDA.write_text('// Generado desde datos/Matriz_App.xlsx con herramientas/matriz_objetivos.py. No editar a mano.\n'
                  'window.MATRIZ_OBJETIVOS = ' + json.dumps(datos, ensure_ascii=False, indent=1) + ';\n', encoding='utf-8')
for t in orden_tipos:
    print(f'{t}: {len(objetivos[t])} objetivos, {sum(len(v) for v in subcategorias.get(t, {}).values())} subcategorías')
print('Variables:', ', '.join(variables) or '—', '| Meses cargados:', ', '.join(sorted(mensual)) or '—')
print('Clasificaciones:', ' · '.join(f'{c}: {"/".join(t)}' for c, t in por_clasificacion.items()) or '—')
for a in avisos:
    print('AVISO:', a)
