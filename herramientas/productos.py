#!/usr/bin/env python3
"""Genera productos.js: catálogo para Productos presentados, Productos pedidos y Muestras del cierre de visita.

Uso:  python3 herramientas/productos.py
Fuentes (en datos/):
- Base_Productos.xlsx: hoja "Base de Productos" (Referencia Interna | Nombre | Etiquetas de producto | Nueva Etiqueta).
  La etiqueta que vale es "Nueva Etiqueta" si tiene algo; si no, "Etiquetas de producto".
  Hoja "Guía de Etiquetas" (Etiqueta | Concepto Estratégico | Mensaje Comercial): sale al abrir la etiqueta en el cierre.
- Productos_Terminados.xlsx: productos terminados (export de Odoo con Referencia Interna, Nombre, Etiquetas de producto).
Categorías del cierre: Nuevo, Foco y Transición-Impulso se despliegan (buscador por código o nombre);
Portafolio y Consultorio van cerradas (solo se marcan).
"""
import collections, json, unicodedata
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
FUENTES = [RAIZ / 'datos' / 'Base_Productos.xlsx', RAIZ / 'datos' / 'Productos_Terminados.xlsx']
DESPLEGABLES = ['Nuevo', 'Foco', 'Transición-Impulso']
CERRADAS = ['Portafolio', 'Consultorio']

def clave(t):
    return unicodedata.normalize('NFD', str(t or '').strip().lower()).encode('ascii', 'ignore').decode()

NOMBRES = {'nuevo': 'Nuevo', 'foco': 'Foco', 'transicion': 'Transición-Impulso', 'transicion-impulso': 'Transición-Impulso',
           'transicion - impulso': 'Transición-Impulso', 'transicion impulso': 'Transición-Impulso', 'portafolio': 'Portafolio',
           'estrategico': 'Estratégico', 'consultorio': 'Consultorio', 'cliente': 'Cliente',
           'a descodificar': 'A descodificar', 'en desarrollo': 'En Desarrollo', 'producto terminado': 'Producto Terminado'}
etiquetas = lambda t: [NOMBRES[clave(x)] for x in str(t or '').replace(';', ',').split(',') if clave(x) in NOMBRES]

productos, vistos, guia = [], set(), {}
for f in FUENTES:
    if not f.exists():
        print('AVISO: no existe', f.relative_to(RAIZ))
        continue
    libro = openpyxl.load_workbook(f, data_only=True)
    for ws in libro.worksheets:
        cab = [clave(c.value) for c in ws[1]]
        if 'concepto estrategico' in cab:   # Guía de etiquetas
            ie, ic, im = cab.index('etiqueta'), cab.index('concepto estrategico'), next(i for i, c in enumerate(cab) if 'mensaje' in c)
            for r in ws.iter_rows(min_row=2, values_only=True):
                for et in etiquetas(r[ie]):
                    guia[et] = {'concepto': ' '.join(str(r[ic] or '').split()), 'mensaje': ' '.join(str(r[im] or '').split())}
            continue
        if 'referencia interna' not in cab:
            continue
        col = lambda n: cab.index(n) if n in cab else None
        ic, inn, ie, inu, ia, ig = col('referencia interna'), col('nombre'), col('etiquetas de producto'), col('nueva etiqueta'), col('activo(a)'), col('grupo de producto')
        for r in ws.iter_rows(min_row=2, values_only=True):
            cod = str(r[ic] or '').strip()
            if not cod or cod in vistos or (ia is not None and r[ia] in (False, 'False', 0)):
                continue
            nueva = r[inu] if inu is not None else None
            vistos.add(cod)
            productos.append({'c': cod, 'n': ' '.join(str(r[inn] or '').split()),
                              'e': etiquetas(nueva if str(nueva or '').strip() else r[ie] if ie is not None else ''),
                              'g': str(r[ig] or '').strip() if ig is not None else ('Producto Terminado' if cod.startswith('PT') else 'Magistral de Pedido')})

productos.sort(key=lambda p: clave(p['n']))
datos = {'desplegables': DESPLEGABLES, 'cerradas': CERRADAS, 'productos': productos, 'guia': guia}
(RAIZ / 'productos.js').write_text('// Generado con herramientas/productos.py desde datos/Base_Productos.xlsx y datos/Productos_Terminados.xlsx. No editar a mano.\n'
                                   'window.CATALOGO = ' + json.dumps(datos, ensure_ascii=False) + ';\n', encoding='utf-8')
cuenta = collections.Counter(e for p in productos for e in p['e'])
for cat in DESPLEGABLES + CERRADAS:
    print(f'{cat}: {cuenta[cat]}')
print('Otras etiquetas:', {k: v for k, v in cuenta.items() if k not in DESPLEGABLES + CERRADAS})
print('Guía de etiquetas:', ', '.join(guia) or '—')
print(f'Total: {len(productos)} productos')
