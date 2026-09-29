#!/usr/bin/env python3
"""Actualiza contactos.json (los clientes de cada zona en la app) desde la Maestra de Contactos de Odoo.

Uso:  python3 herramientas/maestra_contactos.py <Maestra_de_Contactos.xlsx> [AAAA-MM-DD]
      La fecha (por defecto hoy, hora Colombia) es desde cuándo rigen los cambios de zona.

Reglas (definidas por Hernán Reyes):
- No se suben las columnas QUITAR ("Nombre" y "Lista de Precios"); el nombre que usa la app es "Nombre Público".
- El equipo "Empleados" no se sube.
- La zona sale del comercial (COMERCIAL_ZONA); si no está ahí, de "Equipo de ventas".
  Lo que dice "Vacante Epithelium" es de Yunelis Caballero (Zona Sur).
- Los cambios de zona rigen desde la fecha de la subida: lo ya visitado se queda con quien lo visitó.
Deja una copia limpia en datos/Maestra_de_Contactos.xlsx y anota los cambios en datos/historial_maestra.csv.
"""
import csv, json, sys, datetime
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
QUITAR = {'Nombre', 'Lista de Precios'}
NO_SUBIR_EQUIPOS = {'Empleados'}
ZONAS = ['Clientes Especiales', 'Zona Norte', 'Zona Sur']
COMERCIAL_ZONA = {
    'Vacante Epithelium': 'Zona Sur',            # es de Yunelis Caballero
    'Yunelis Caballero': 'Zona Sur',
    'Lizeth Geraldine Ramos Guerrero': 'Zona Norte',
    'Jennifer Andrea Herrera': 'Clientes Especiales',
}

entrada = Path(sys.argv[1])
fecha = sys.argv[2] if len(sys.argv) > 2 else (datetime.datetime.utcnow() - datetime.timedelta(hours=5)).date().isoformat()
ws = openpyxl.load_workbook(entrada, data_only=True).active
cab = [str(c.value or '').strip() for c in ws[1]]
col = {n: i for i, n in enumerate(cab)}
for n in ('Nombre Público', 'Ciudad', 'Comercial', 'Equipo de ventas', 'Etiquetas'):
    if n not in col:
        sys.exit(f'Falta la columna "{n}" en la maestra')

def zona_de(comercial, equipo):
    for nombre, zona in COMERCIAL_ZONA.items():
        if comercial and nombre.lower() in comercial.lower():
            return zona
    return equipo if equipo in ZONAS else None

nuevo, sin_zona = {z: [] for z in ZONAS}, []
filas_limpias = []
for fila in ws.iter_rows(min_row=2, values_only=True):
    if not any(fila):
        continue
    val = lambda n: str(fila[col[n]]).strip() if fila[col[n]] not in (None, False) else ''
    equipo, comercial = val('Equipo de ventas'), val('Comercial')
    if equipo in NO_SUBIR_EQUIPOS:
        continue
    nombre = val('Nombre Público')
    if not nombre:
        continue
    zona = zona_de(comercial, equipo)
    if not zona:
        sin_zona.append(nombre)
        continue
    # n nombre, c ciudad, e etiquetas, t título, p provincia, f cliente para facturar, cl clasificación, ca categoría, pz plazo de pago
    facturar = fila[col['Cliente para Facturar']] if 'Cliente para Facturar' in col else None
    extra = {'t': val('Título') if 'Título' in col else '', 'p': val('Provincia') if 'Provincia' in col else '',
             'f': facturar in (True, 'True', 'true', 1, 'Sí', 'Si'),
             'cl': val('Categoría de cliente/Clasificación') if 'Categoría de cliente/Clasificación' in col else '',
             'ca': val('Categoría de cliente') if 'Categoría de cliente' in col else '',
             'pz': val('Plazos de Pago') if 'Plazos de Pago' in col else ''}
    nuevo[zona].append({'n': nombre, 'c': val('Ciudad'), 'e': val('Etiquetas'), **{k: v for k, v in extra.items() if v not in ('', None)}})
    filas_limpias.append([fila[i] for i, n in enumerate(cab) if n not in QUITAR] + [zona])
for z in ZONAS:
    nuevo[z].sort(key=lambda x: x['n'].lower())

anterior = json.loads((RAIZ / 'contactos.json').read_text(encoding='utf-8'))
zona_ant = {x['n']: z for z, l in anterior.items() for x in l}
zona_new = {x['n']: z for z, l in nuevo.items() for x in l}
cambios = [(n, 'nuevo', '', z) for n, z in zona_new.items() if n not in zona_ant] \
    + [(n, 'retirado', z, '') for n, z in zona_ant.items() if n not in zona_new] \
    + [(n, 'cambio de zona', zona_ant[n], z) for n, z in zona_new.items() if n in zona_ant and zona_ant[n] != z]

(RAIZ / 'contactos.json').write_text(json.dumps(nuevo, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
# Copia limpia (sin las columnas que no se suben) y sin cuadrícula
limpio = openpyxl.Workbook(); h = limpio.active; h.title = 'Maestra'
h.append([n for n in cab if n not in QUITAR] + ['Zona en la app'])
for f in filas_limpias:
    h.append(f)
h.sheet_view.showGridLines = False
h.freeze_panes = 'A2'
limpio.save(RAIZ / 'datos' / 'Maestra_de_Contactos.xlsx')
hist = RAIZ / 'datos' / 'historial_maestra.csv'
nuevo_archivo = not hist.exists()
with hist.open('a', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    if nuevo_archivo:
        w.writerow(['Vigente desde', 'Contacto', 'Cambio', 'Zona anterior', 'Zona nueva'])
    for c in cambios:
        w.writerow([fecha] + list(c))
print('Contactos por zona:', {z: len(l) for z, l in nuevo.items()})
print(f'Cambios (vigentes desde {fecha}):', len(cambios))
for c in cambios:
    print('  -', ' · '.join(x for x in c if x))
if sin_zona:
    print('SIN ZONA (no se subieron):', ', '.join(sin_zona))
