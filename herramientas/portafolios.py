#!/usr/bin/env python3
"""Genera portafolios.js: portafolio exclusivo de cada cliente para la Maestra de Clientes.

Uso:  python3 herramientas/portafolios.py [carpeta del Vademécum]   (por defecto ../vademecum-epithelium)
- Productos: los de datos/Base_Productos.xlsx con etiqueta Cliente y categoría "Magistral de Pedido / <cliente>".
  Las categorías generales (Acné, Despigmentantes...) no son portafolio exclusivo y no salen.
- Ficha (componentes, indicación y dosis): del Vademécum, que es la referencia (portafolios/*.json y data.json).
- Cliente de la categoría -> contactos de la Maestra (contactos.json): por nombre igual, por sedes ("Cliente, Sede")
  y por la tabla MAPA para los que están escritos distinto. Lo que no cruce se avisa al final.
"""
import json, re, sys, unicodedata
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
VADE = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ.parent / 'vademecum-epithelium'

def clave(t):
    t = unicodedata.normalize('NFD', str(t or '').strip().lower()).encode('ascii', 'ignore').decode()
    return ' '.join(re.sub(r'[^a-z0-9@&]+', ' ', t).split())

# Cliente (como está en la categoría) -> contactos de la Maestra cuando el nombre no es igual
MAPA = {
    'Darwin Alejandro Lazo': ['Darwin Alejandro Lazo Sanchez'],
    'Dermaskin Cuidado para La Piel - Dra Claudia Milena Arango': ['Claudia Milena Arango Atehortua'],
    'Dermatologia y Medicina Estetica- Dra. Erika Molina/Deladieth Serrano Rincon': ['Erika Johanna Molina Jaime', 'Deladieth Serrano Rincon'],
    'Dr. Carlos Fernando Quintero': ['Carlos Fernando Quintero Baute'],
    'Dr. Felix Duran Gonzalez': ['Felix Duran Gonzalez'],
    'Dr. Juan Sebastian Suarez': ['Juan Sebastian Suarez Arango'],
    'Dra. Carol Mesa Dermatologa': ['Cm Dra Carol Mesa Dermatologa'],
    'Dra. Diana Elizabeth Reyes Bocanegra': ['Diana Elizabeth Reyes Bocanegra'],
    'Dra. Natalia Hernandez Mantilla': ['Natalia Hernandez Mantilla'],
    'Dra. Tatiana Leal S.A.S.': ['Dra. Tatiana Leal Sandoval', 'Tatiana Alexandra Leal Sandoval'],
    'Drogueria Dermatos - Dr. Andres Mauricio Ardila': ['Andres Mauricio Ardila Rojas'],
    'Epidermica S.A.S.- Dra. Alejandra Zuluaga': ['Alejandra Zuluaga Cardona'],
    'Es Vital S.A.S.': ['Es Vital Center S.A.S.'],
    'Farmadescuentos Co S.A.S - Pharmags Laboratories S.A.S.': ['Farmadescuentos Co S.A.S'],
    'Farmapiel Garzon - Carolina Andrea Murcia Alvarez': ['Carolina Andrea Murcia Alvarez'],
    'GG & Piel S.A.S - Dr. Giovanny Antonio Guerra': ['Giovanny Antonio Guerra Jimenez'],
    'Gema Clinic - Dra. Lizeth Daniela Caro Suarez': ['Lizeth Daniela Caro Suarez'],
    'Guapa Glow - Sandra Carolina Moreno': ['Sandra Carolina Moreno Ferrer'],
    'Hair and Skin Evolution - Carolina Palacio S.A.S.': ['Carolina Palacio S.A.S.'],
    'Niks Leip - Maria de Jesus Acosta Correa': ['Maria De Jesus Correa Acosta'],
    'PielMedic Cabello y Piel - Antony Adrian Aguilar Erazo': ['Pielmedic, Antony Adrian Aguilar Erazo'],
    'Plenufarma - Dr. Jesus Enrique Vasquez': ['Jesus Enrique Vasquez Montero'],
}

# ---- Productos con etiqueta Cliente y categoría con el nombre del cliente
NOMBRES_ETQ = lambda t: [clave(x) for x in str(t or '').replace(';', ',').split(',')]
ws = openpyxl.load_workbook(RAIZ / 'datos' / 'Base_Productos.xlsx', data_only=True)['Base de Productos']
cab = [clave(c.value) for c in ws[1]]
ic, inn, icat, ie, inu = (cab.index(n) for n in ('referencia interna', 'nombre', 'categoria del producto', 'etiquetas de producto', 'nueva etiqueta'))
por_cliente = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    et = r[inu] if str(r[inu] or '').strip() else r[ie]
    if 'cliente' not in NOMBRES_ETQ(et) or ' / ' not in str(r[icat] or ''):
        continue
    por_cliente.setdefault(str(r[icat]).split(' / ', 1)[1].strip(), []).append((str(r[ic]).strip(), ' '.join(str(r[inn] or '').split())))

# ---- Fichas del Vademécum: primero la del portafolio del mismo cliente, si no cualquiera con ese código
fichas, fichas_cli = {}, {}
datos = json.loads((VADE / 'data.json').read_text())
for f in sorted((VADE / 'portafolios').glob('*.json')):
    d = json.loads(f.read_text())
    for p in d['productos']:
        fichas_cli[(clave(d['cliente']), p['Referencia Interna'])] = p
        fichas.setdefault(p['Referencia Interna'], p)
for p in (datos if isinstance(datos, list) else datos.get('productos', [])):
    fichas.setdefault(p['Referencia Interna'], p)

def ficha(cliente, cod):
    k = clave(cliente)
    p = fichas_cli.get((k, cod)) or next((v for (c, x), v in fichas_cli.items() if x == cod and (c.startswith(k) or k.startswith(c))), None) or fichas.get(cod) or {}
    return {'comp': p.get('Componentes') or '', 'ind': p.get('Indicación') or '', 'dosis': p.get('Dosis Recomendada') or ''}

# ---- Clientes de la Maestra
contactos = json.loads((RAIZ / 'contactos.json').read_text())
nombres = sorted({c['n'] for l in contactos.values() for c in l})
generales = []
salida, contactos_de = {}, {}
for cliente, prods in sorted(por_cliente.items(), key=lambda x: clave(x[0])):
    k = clave(cliente)
    # Las categorías generales (no son un cliente) no cruzan con la Maestra
    # Nombre igual, sus sedes ("Cliente, Sede") o la empresa de "Empresa - Dueño"
    empresa = clave(re.split(r'\s*-\s+', cliente)[0])
    destino = [n for n in nombres if clave(n) == k or clave(n.split(',')[0]) in (k, empresa)]
    destino += [n for x in MAPA.get(cliente, []) for n in nombres if n == x or n.startswith(x + ',')]
    destino = sorted(set(destino))
    if not destino:
        generales.append((cliente, len(prods)))
        continue
    salida[cliente] = [{'c': c, 'n': n, **ficha(cliente, c)} for c, n in sorted(prods, key=lambda p: clave(p[1]))]
    for n in destino:
        contactos_de[clave(n)] = cliente

sin_ficha = [p['c'] for l in salida.values() for p in l if not (p['ind'] or p['dosis'])]
(RAIZ / 'portafolios.js').write_text(
    '// Generado con herramientas/portafolios.py desde datos/Base_Productos.xlsx (productos) y el Vademécum (fichas). No editar a mano.\n'
    'window.PORTAFOLIOS = ' + json.dumps({'clientes': contactos_de, 'portafolios': salida}, ensure_ascii=False) + ';\n')
print(f'{len(salida)} clientes con portafolio exclusivo, {sum(map(len, salida.values()))} productos, {len(contactos_de)} contactos de la Maestra')
if generales:
    print('Sin cliente en la Maestra (no salen):', generales)
if sin_ficha:
    print('Sin indicación ni dosis en el Vademécum:', sin_ficha)
