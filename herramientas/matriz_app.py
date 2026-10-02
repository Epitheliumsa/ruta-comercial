#!/usr/bin/env python3
"""Arma o refresca la matriz completa de la app en un solo Excel: datos/Matriz_App.xlsx.

Uso:  python3 herramientas/matriz_app.py [Excel entregado por el usuario]

Hojas:
- Índice: qué hay en cada hoja y el resumen de lo que carga la app.
- Visitas / Trabajo interno / Mensual: objetivos y subcategorías (ver herramientas/matriz_objetivos.py).
- Tipo de visita: en qué tipo de visita sale cada clasificación de cliente (X). Los conteos de clientes y
  etiquetas se recalculan con contactos.json; las X se conservan. Clasificaciones nuevas salen sin X (en amarillo).
Después de armarla, corre matriz_objetivos.py para regenerar objetivos.js.
"""
import collections, datetime, json, subprocess, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / 'datos' / 'Matriz_App.xlsx'
VIEJA_OBJ = RAIZ / 'datos' / 'Matriz_objetivos_subcategorias.xlsx'
VIEJA_CLASIF = RAIZ / 'datos' / 'Matriz_tipo_visita_clasificacion.xlsx'
TIPOS = ['Visita Médica', 'Visita Cliente', 'Punto de Venta']
HOJA_TIPO = 'Tipo de visita'
VERDE = '006B4F'

entrada = Path(sys.argv[1]) if len(sys.argv) > 1 else SALIDA if SALIDA.exists() else VIEJA_OBJ
libro = openpyxl.load_workbook(entrada)

# X por clasificación: de la hoja del libro o, si no la tiene, del archivo viejo de clasificaciones
def leer_x(ws):
    cab = [str(c.value or '').strip() for c in ws[4]]
    cols = {i: {'Visita Comercial': 'Visita Cliente'}.get(n, n) for i, n in enumerate(cab) if n in TIPOS or n == 'Visita Comercial'}   # acepta el nombre viejo
    x = {}
    for fila in ws.iter_rows(min_row=5, values_only=True):
        cl = str(fila[0] or '').strip()
        if cl.isdigit():
            x[cl] = [n for i, n in cols.items() if str(fila[i] or '').strip()]
    return x
if HOJA_TIPO in libro.sheetnames:
    marcas = leer_x(libro[HOJA_TIPO])
    del libro[HOJA_TIPO]
else:
    marcas = leer_x(openpyxl.load_workbook(VIEJA_CLASIF).active) if VIEJA_CLASIF.exists() else {}

# Clientes de la Maestra (contactos.json) por clasificación
contactos = json.loads((RAIZ / 'contactos.json').read_text(encoding='utf-8'))
cat, n, et = {}, collections.Counter(), collections.defaultdict(collections.Counter)
for zona in contactos.values():
    for c in zona:
        cl = str(c.get('cl') or '').strip()
        if not cl:
            continue
        cat[cl] = c.get('ca', '') or cat.get(cl, '')
        n[cl] += 1
        et[cl][c.get('e', '')] += 1
for cl in marcas:
    cat.setdefault(cl, '')

ws = libro.create_sheet(HOJA_TIPO)
ws.sheet_view.showGridLines = False
ws['A1'] = 'Tipo de visita según la clasificación del cliente'
ws['A1'].font = Font(bold=True, size=14, color=VERDE)
ws['A2'] = ('X en los tipos de visita donde sale cada clasificación. Con X en Visita Médica y Visita Cliente (20 y 21) el vendedor '
            'marca una, otra o ambas; con las dos marcadas se usan los objetivos de "Visita Médica Comercial" (hoja Visitas).')
ws['A2'].font = Font(italic=True, color='666666')
ws.append([])
ws.append(['Clasificación', 'Categoría', 'Clientes', 'Etiquetas en la maestra'] + TIPOS)
amarillo = PatternFill('solid', fgColor='FFF7D6')
aviso = PatternFill('solid', fgColor='FDE68A')
borde = Side(style='thin', color='E5D9A8')
for cl in sorted(cat, key=int):
    ws.append([int(cl), cat[cl], n[cl], ' · '.join(f'{e} ({k})' for e, k in et[cl].most_common()) or '—']
              + ['X' if t in marcas.get(cl, []) else None for t in TIPOS])
fin = ws.max_row
t = Table(displayName='TipoDeVisita', ref=f'A4:G{fin}')
t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium7', showRowStripes=True)
ws.add_table(t)
dv = DataValidation(type='list', formula1='"X"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'E5:G{fin + 20}')
for i, w in enumerate([14, 58, 10, 60, 16, 17, 16], 1):
    ws.column_dimensions[chr(64 + i)].width = w
for r in range(4, fin + 1):
    sin_x = r > 4 and not any(ws.cell(r, c).value for c in (5, 6, 7))
    for c in range(1, 8):
        celda = ws.cell(r, c)
        celda.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if c in (1, 3, 5, 6, 7) else 'left')
        if r > 4 and c >= 5:
            celda.fill = aviso if sin_x else amarillo
            celda.font = Font(bold=True, size=12, color=VERDE)
            celda.border = Border(left=borde, right=borde, top=borde, bottom=borde)
ws.freeze_panes = 'C5'
ws.cell(fin + 2, 1, f'Total clientes en la maestra: {sum(n.values())} · Amarillo fuerte = clasificación sin X (sale en todos los tipos)').font = Font(italic=True, color='666666')

# Índice
if 'Índice' in libro.sheetnames:
    del libro['Índice']
ix = libro.create_sheet('Índice', 0)
ix.sheet_view.showGridLines = False
ix['A1'] = 'Matriz de Epithelium Visita'
ix['A1'].font = Font(bold=True, size=16, color=VERDE)
ix['A2'] = f'Actualizada el {datetime.date.today().strftime("%d/%m/%Y")} · es lo que la app tiene cargado hoy'
ix['A2'].font = Font(italic=True, color='666666')
filas = [('Hoja', 'Qué contiene'),
         ('Visitas', 'Objetivos (fila verde, ✓) y subcategorías (X) por tipo de visita: Visita Médica, Visita Médica Comercial '
                     '(20 y 21 con las dos marcadas), Visita Cliente, Punto de Venta y Contacto nuevo.'),
         ('Trabajo interno', 'Objetivos y subcategorías de Oficina, Fuera de la Oficina, Planeación Mes y Mercadeo (solo Coordinadora Comercial).'),
         ('Mensual', 'Parrilla Promocional y Actividades de cada mes (subcategorías variables).'),
         (HOJA_TIPO, 'En qué tipo de visita sale cada clasificación de cliente de la Maestra.')]
for i, (a, b) in enumerate(filas, 4):
    ix.cell(i, 1, a); ix.cell(i, 2, b)
    ix.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')
    if i == 4:
        for c in (1, 2):
            ix.cell(i, c).font = Font(bold=True, color='FFFFFF')
            ix.cell(i, c).fill = PatternFill('solid', fgColor=VERDE)
    else:
        ix.cell(i, 1).font = Font(bold=True, color=VERDE)
ix['A11'] = 'Reglas'
ix['A11'].font = Font(bold=True, color=VERDE)
reglas = ['Nombres en nombre propio (NOMPROPIO), con de/del/la/y en minúscula: "Chequeo de Precios".',
          'Objetivos en orden alfabético (salvo Planeación Mes). Subcategorías: el número en la columna A es su orden.',
          'Clasificaciones 20 y 21: al programar se marca Visita Médica, Visita Médica Comercial o ambas (lista unida).',
          'No cambies los encabezados (fila 4) ni el nombre de las hojas.']
for i, r in enumerate(reglas, 12):
    ix.cell(i, 1, '•'); ix.cell(i, 2, r); ix.cell(i, 2).alignment = Alignment(wrap_text=True)
ix.column_dimensions['A'].width = 18
ix.column_dimensions['B'].width = 110

for h in libro.worksheets:
    h.sheet_view.showGridLines = False
libro.active = 0
libro.save(SALIDA)
print('Matriz:', SALIDA.relative_to(RAIZ), '| hojas:', ', '.join(libro.sheetnames))
sin = [cl for cl in cat if not marcas.get(cl)]
if sin:
    print('AVISO: clasificaciones sin X (salen en todos los tipos):', ', '.join(sorted(sin, key=int)))
subprocess.run([sys.executable, str(RAIZ / 'herramientas' / 'matriz_objetivos.py')], check=True)
