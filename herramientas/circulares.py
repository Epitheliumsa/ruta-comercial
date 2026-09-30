#!/usr/bin/env python3
"""Genera circulares.js: circulares de Ventas/Mercadeo para Actividades-Circulares y el objetivo "Actividades".

Uso:  python3 herramientas/circulares.py [Excel]
Fuente: datos/Circulares.xlsx, hoja "Circulares" (fila 1 = encabezados):
  Circular | Tipo | Nombre de la actividad | Grupo | Dirigida a (canal/cliente) | Producto | Fecha inicio | Fecha fin |
  Estado | Días vencida / restantes | Objetivo | Resumen de la actividad | Observación | Enlace PDF
Estado y Días los calcula la app cada día (las columnas del Excel se ignoran). Fecha fin vacía o "Indefinido" = sin fin.
"Dirigida a" dice en qué visitas sale la circular (objetivo "Actividades"):
- Canales / subcanales con números (ej: "Canales 20, 21, 22") -> clientes con esa clasificación.
- Nombres de clientes (ej: "Bella Piel S.A.S.") -> los clientes de la Maestra que se llaman así.
- "(no aplica X)" -> quita esos clientes.   "Todos" -> todos los clientes.   "Médicos…" sin números -> 20, 21 y 22.
- Internas (Fuerza de Ventas, Equipo, Coordinadora, Canal 80 colaboradores) -> no salen en visitas.
Si se pasa otro Excel, se copia a datos/Circulares.xlsx.
"""
import datetime, json, re, shutil, sys, unicodedata
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / 'datos' / 'Circulares.xlsx'
SALIDA = RAIZ / 'circulares.js'
INTERNAS = ['fuerza de ventas', 'equipo visita', 'coordinador', 'colaboradores']
QUITAR = ['s.a.s.', 's.a.s', 'sas', 's.a.', 'ltda', 'de colombia']

def clave(t):
    t = unicodedata.normalize('NFD', str(t or '').lower()).encode('ascii', 'ignore').decode()
    for q in QUITAR:
        t = t.replace(q, ' ')
    return ' '.join(re.sub(r'[^a-z0-9 ]', ' ', t).split())

def fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    t = str(v or '').strip()
    return t if re.fullmatch(r'\d{4}-\d{2}-\d{2}', t) else ''

if len(sys.argv) > 1 and Path(sys.argv[1]).resolve() != FUENTE.resolve():
    shutil.copy(sys.argv[1], FUENTE)

contactos = json.loads((RAIZ / 'contactos.json').read_text(encoding='utf-8'))
clientes = sorted({c['n'] for zona in contactos.values() for c in zona})

def clientes_de(texto):
    """Clientes de la Maestra que nombra el texto (por partes separadas por coma, "y" o paréntesis)."""
    partes = [clave(p) for p in re.split(r'[,;()]| y ', texto)]
    return sorted({n for n in clientes for p in partes if len(p) >= 4 and (p in clave(n) or clave(n) in p)})

ws = openpyxl.load_workbook(FUENTE, data_only=True)['Circulares']
cab = [clave(c.value) for c in ws[1]]
col = lambda nombre: next((i for i, c in enumerate(cab) if c.startswith(nombre)), None)
I = {k: col(k) for k in ['circular', 'tipo', 'nombre', 'grupo', 'dirigida', 'producto', 'fecha inicio', 'fecha fin',
                          'objetivo', 'resumen', 'observacion', 'enlace']}
dato = lambda fila, k: fila[I[k]] if I[k] is not None and I[k] < len(fila) else None

circulares, avisos = [], []
for fila in ws.iter_rows(min_row=2, values_only=True):
    codigo = str(dato(fila, 'circular') or '').strip()
    if not codigo:
        continue
    dirigida = str(dato(fila, 'dirigida') or '').strip()
    d = clave(dirigida)
    principal, _, excluye = dirigida.partition('no aplica')
    canales = sorted(set(re.findall(r'\b(\d{2})\b', principal))) if re.search(r'canal', d) else []
    interna = any(x in d for x in INTERNAS) or canales == ['80']
    todos = d == 'todos'
    if not canales and not interna and not todos and 'medico' in d:
        canales = ['20', '21', '22']
    nombres = [] if interna or todos or canales else clientes_de(principal)
    excluidos = clientes_de(excluye) if excluye else []
    if not (interna or todos or canales or nombres):
        avisos.append(f'{codigo}: "{dirigida}" no coincide con ningún cliente de la Maestra (no saldrá en visitas)')
    c = {
        'c': codigo, 'tipo': str(dato(fila, 'tipo') or '').strip(), 'nombre': str(dato(fila, 'nombre') or '').strip(),
        'grupo': str(dato(fila, 'grupo') or '').strip(), 'dirigida': dirigida, 'producto': str(dato(fila, 'producto') or '').strip(),
        'ini': fecha(dato(fila, 'fecha inicio')), 'fin': fecha(dato(fila, 'fecha fin')),
        'objetivo': str(dato(fila, 'objetivo') or '').strip(), 'resumen': str(dato(fila, 'resumen') or '').strip(),
        'obs': str(dato(fila, 'observacion') or '').strip(), 'pdf': str(dato(fila, 'enlace') or '').strip(),
        'canales': canales, 'clientes': nombres, 'excluidos': excluidos, 'todos': todos, 'interna': interna,
    }
    if not c['ini']:
        avisos.append(f'{codigo}: sin fecha de inicio')
    circulares.append(c)

SALIDA.write_text('// Generado desde datos/Circulares.xlsx con herramientas/circulares.py. No editar a mano.\n'
                  'window.CIRCULARES = ' + json.dumps(circulares, ensure_ascii=False, indent=1) + ';\n', encoding='utf-8')
print(f'{len(circulares)} circulares -> {SALIDA.name}')
for c in circulares:
    a = 'interna' if c['interna'] else 'todos' if c['todos'] else ('canales ' + ', '.join(c['canales'])) if c['canales'] else f"{len(c['clientes'])} clientes: " + ', '.join(c['clientes'][:4]) + ('…' if len(c['clientes']) > 4 else '')
    print(f"  {c['c']} · {c['nombre'][:45]:45} · {c['ini']} a {c['fin'] or 'sin fin'} · {a}" + (f" · sin {len(c['excluidos'])} clientes" if c['excluidos'] else ''))
for a in avisos:
    print('AVISO:', a)
